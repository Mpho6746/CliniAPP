"""
CliniApp — clinic management system.

Run with:
    pip install -r requirements.txt
    python app.py

Then open http://127.0.0.1:5000
"""

import calendar
import csv
import io
import os
from datetime import date, datetime, timedelta

from flask import Flask, Response, flash, redirect, render_template, request, session, url_for
from openpyxl import Workbook

from charts import bar_chart, line_chart, month_grid, pie_chart

from models import (
    ACCESS_LEVEL_LABELS,
    ACCESS_LEVELS,
    DEFAULT_ACCESS_LEVEL_BY_ROLE,
    STAFF_GENDER_OPTIONS,
    STAFF_LABELS,
    STAFF_ROLES,
    STAFF_SHIFT_OPTIONS,
    Staff,
    User,
    add_consultation,
    add_lab_result,
    add_prescription,
    age_gender_breakdown,
    appointment_counts_for_month,
    appointments_on_date,
    count_consultations_by_doctor,
    count_lab_results_by_technician,
    count_prescriptions_by_doctor,
    create_patient,
    create_staff,
    create_ticket,
    create_user,
    db,
    DEFAULT_DASHBOARD_WIDGETS,
    delete_patient,
    demographics_summary,
    employee_number_taken,
    find_potential_duplicate_patients,
    find_staff_by_employee_number,
    find_user_by_hospital_number,
    get_dashboard_layout,
    get_patient,
    get_settings,
    get_ticket,
    has_access,
    is_overdue,
    last_visit,
    list_all_tickets,
    list_patients,
    list_tickets_by_submitter,
    mark_patient_notes_read,
    new_patients_per_day,
    new_patients_since,
    new_registrations_by_month,
    open_ticket_count,
    overdue_appointment_count,
    patient_consultations,
    patient_growth,
    patient_lab_results,
    patient_missing_fields,
    patient_prescriptions,
    PATIENT_STATUSES,
    patient_number_taken,
    recent_activity,
    recent_patients,
    recent_staff,
    resolve_ticket,
    save_dashboard_order,
    search_staff,
    staff_counts_by_role,
    theme_for_role,
    toggle_dashboard_widget,
    total_patient_count,
    total_staff_count,
    unread_notes_by_patient,
    unread_notes_count,
    update_announcement,
    update_clinic_profile,
    update_department_themes,
    update_patient,
    update_pin_policy,
)

SEARCHABLE_PAGES = {
    "admin": [
        ("Overview", "workspace"),
        ("Patients", "admin_patients"),
        ("Add Staff", "staff_register"),
        ("Report", "admin_report"),
        ("Settings", "admin_settings"),
    ],
    "doctor": [
        ("Overview", "workspace"),
        ("Patients", "doctor_patients"),
    ],
    "nurse": [("Overview", "workspace")],
    "pharmacist": [("Overview", "workspace")],
    "receptionist": [
        ("Overview", "workspace"),
        ("Patients", "receptionist_patients"),
    ],
    "lab_technician": [
        ("Overview", "workspace"),
        ("Patients", "lab_patients"),
    ],
}

FAQ_ITEMS = [
    ("How do I register a new patient?", "Admin > Patients > Register patient. A patient number is generated automatically."),
    ("How do I add a new staff member?", "Admin > Add Staff. Choose their role, then set an employee number and PIN."),
    ("How do I change the PIN policy?", "Admin > Settings > Staff PIN policy. This applies to staff added from then on."),
    ("How do I switch light/dark mode?", "Click the sun/moon icon in the top bar to toggle your own department's theme."),
    ("How do I search for a patient or staff member?", "Use the search box in the top bar — it also matches pages like Report or Settings."),
    ("I forgot my employee number or PIN.", "There's no self-service recovery yet — submit a support ticket below and an Admin will help."),
]

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)
app.secret_key = "dev-secret-key-change-me"  # replace with a real secret in production
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(BASE_DIR, "clinic.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db.init_app(app)
with app.app_context():
    db.create_all()


def current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    return db.session.get(User, user_id)


@app.context_processor
def inject_theme():
    role, _ = current_staff()
    return {"theme": theme_for_role(role), "theme_role": role}


@app.context_processor
def inject_notifications():
    _, staff = current_staff()
    return {"notifications": recent_activity(8) if staff else []}


@app.context_processor
def inject_announcement():
    _, staff = current_staff()
    return {"announcement": get_settings().announcement if staff else ""}


@app.context_processor
def inject_has_access():
    return {"has_access": has_access}


def current_staff():
    """The staff role/record for the second login step, if completed."""
    staff_id = session.get("staff_id")
    if not staff_id:
        return None, None
    staff = db.session.get(Staff, staff_id)
    return (staff.role, staff) if staff else (None, None)


# -------------------- Auth --------------------

@app.route("/")
def home():
    if current_user():
        return redirect(url_for("dashboard"))
    return redirect(url_for("clinic_login"))


@app.route("/clinic-login", methods=["GET", "POST"])
def clinic_login():
    if current_user():
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        hospital_number = request.form.get("hospital_number", "").strip()
        password = request.form.get("password", "")

        user = find_user_by_hospital_number(hospital_number)
        if user is None or not user.check_password(password):
            flash("Invalid hospital number or password.", "error")
            return redirect(url_for("clinic_login"))

        session.clear()
        session["user_id"] = user.id
        flash(f"Welcome back, {user.hospital_name}.", "success")
        return redirect(url_for("dashboard"))

    return render_template("clinic_login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user():
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        hospital_name = request.form.get("hospital_name", "").strip()
        hospital_number = request.form.get("hospital_number", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")

        errors = []
        if not hospital_name:
            errors.append("Enter your hospital name.")
        if not hospital_number:
            errors.append("Enter your hospital number.")
        elif find_user_by_hospital_number(hospital_number):
            errors.append("That hospital number is already registered.")
        if len(password) < 6:
            errors.append("Password must be at least 6 characters.")
        if password != confirm:
            errors.append("Passwords don't match.")

        if errors:
            for e in errors:
                flash(e, "error")
            return redirect(url_for("register"))

        create_user(hospital_number, password, hospital_name)
        flash("Account created. You can now log in.", "success")
        return redirect(url_for("clinic_login"))

    return render_template("register.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("clinic_login"))


# -------------------- Staff login (second step, after the welcome page) --------------------

@app.route("/staff", methods=["GET", "POST"])
def staff_login():
    if not current_user():
        return redirect(url_for("clinic_login"))

    role, staff = current_staff()
    if staff:
        return redirect(url_for("workspace"))

    if request.method == "POST":
        employee_number = request.form.get("employee_number", "").strip()
        pin = request.form.get("pin", "").strip()

        role, staff = find_staff_by_employee_number(employee_number)
        if staff is None or not staff.check_pin(pin):
            flash("Invalid employee number or PIN.", "error")
            return redirect(url_for("staff_login"))

        session["staff_role"] = role
        session["staff_id"] = staff.id
        flash(f"Signed in as {STAFF_LABELS[role]} {staff.full_name}.", "success")
        return redirect(url_for("workspace"))

    return render_template("staff_login.html", bootstrap=total_staff_count() == 0)


@app.route("/staff/register", methods=["GET", "POST"])
def staff_register():
    if not current_user():
        return redirect(url_for("clinic_login"))

    # Bootstrap case: no staff exist anywhere yet, so there's no admin who
    # could add one. Anyone with a hospital login may set up the first
    # account, and it must be an Admin so they can add everyone else.
    bootstrap = total_staff_count() == 0
    admin_staff = None
    if bootstrap:
        allowed_roles = {"admin": STAFF_LABELS["admin"]}
    else:
        role, staff = current_staff()
        if not staff:
            return redirect(url_for("staff_login"))
        if role != "admin":
            flash("Only admins can add staff members.", "error")
            return redirect(url_for("workspace"))
        if not has_access(staff, "admin_ops", "write"):
            flash("Your account doesn't have permission to add staff members.", "error")
            return redirect(url_for("workspace"))
        allowed_roles = STAFF_LABELS
        admin_staff = staff

    pin_length = get_settings().pin_length

    if request.method == "POST":
        role = request.form.get("role", "")
        full_name = request.form.get("full_name", "").strip()
        employee_number = request.form.get("employee_number", "").strip()
        pin = request.form.get("pin", "").strip()
        confirm = request.form.get("confirm", "").strip()

        email = request.form.get("email", "").strip()
        phone_work = request.form.get("phone_work", "").strip()
        phone_mobile = request.form.get("phone_mobile", "").strip()
        gender = request.form.get("gender", "").strip()
        national_id = request.form.get("national_id", "").strip()
        home_address = request.form.get("home_address", "").strip()

        specialization = request.form.get("specialization", "").strip()
        qualification = request.form.get("qualification", "").strip()
        license_number = request.form.get("license_number", "").strip()
        shift_preference = request.form.get("shift_preference", "").strip()

        access_level = request.form.get("access_level", "").strip()

        errors = []
        if role not in allowed_roles:
            errors.append("Choose a valid role.")
        if bootstrap:
            access_level = "full"
        elif access_level not in ACCESS_LEVELS:
            errors.append("Choose a valid access level.")
        if not full_name:
            errors.append("Enter your full name.")
        if not employee_number.isdigit() or len(employee_number) != 6:
            errors.append("Employee number must be exactly 6 digits.")
        elif employee_number_taken(employee_number):
            errors.append("That employee number is already registered.")
        if not pin.isdigit() or len(pin) != pin_length:
            errors.append(f"PIN must be exactly {pin_length} digits.")
        elif pin != confirm:
            errors.append("PINs don't match.")

        date_of_birth_raw = request.form.get("date_of_birth", "").strip()
        date_of_birth = None
        if date_of_birth_raw:
            try:
                date_of_birth = datetime.strptime(date_of_birth_raw, "%Y-%m-%d").date()
            except ValueError:
                errors.append("Date of birth is not a valid date.")

        if gender and gender not in STAFF_GENDER_OPTIONS:
            errors.append("Choose a valid gender.")

        years_of_experience_raw = request.form.get("years_of_experience", "").strip()
        years_of_experience = None
        if years_of_experience_raw:
            if years_of_experience_raw.isdigit() and 0 <= int(years_of_experience_raw) <= 70:
                years_of_experience = int(years_of_experience_raw)
            else:
                errors.append("Years of experience must be a number between 0 and 70.")

        date_joined_raw = request.form.get("date_joined", "").strip()
        date_joined = None
        if date_joined_raw:
            try:
                date_joined = datetime.strptime(date_joined_raw, "%Y-%m-%d").date()
            except ValueError:
                errors.append("Date of joining is not a valid date.")

        if shift_preference and shift_preference not in STAFF_SHIFT_OPTIONS:
            errors.append("Choose a valid shift preference.")

        consultation_fee_raw = request.form.get("consultation_fee", "").strip()
        consultation_fee = None
        if role == "doctor" and consultation_fee_raw:
            try:
                consultation_fee = float(consultation_fee_raw)
                if consultation_fee < 0:
                    raise ValueError
            except ValueError:
                errors.append("Consultation fee must be a positive number.")

        if errors:
            for e in errors:
                flash(e, "error")
            return redirect(url_for("staff_register"))

        create_staff(
            role,
            employee_number,
            pin,
            full_name,
            access_level=access_level,
            email=email,
            phone_work=phone_work,
            phone_mobile=phone_mobile,
            date_of_birth=date_of_birth,
            gender=gender,
            national_id=national_id,
            home_address=home_address,
            specialization=specialization,
            qualification=qualification,
            years_of_experience=years_of_experience,
            license_number=license_number,
            date_joined=date_joined,
            shift_preference=shift_preference,
            consultation_fee=consultation_fee,
        )
        flash("Staff account created.", "success")
        return redirect(url_for("staff_login") if bootstrap else url_for("workspace"))

    return render_template(
        "staff_register.html",
        roles=allowed_roles,
        bootstrap=bootstrap,
        pin_length=pin_length,
        staff=admin_staff,
        gender_options=STAFF_GENDER_OPTIONS,
        shift_options=STAFF_SHIFT_OPTIONS,
        access_levels=ACCESS_LEVELS,
        access_level_labels=ACCESS_LEVEL_LABELS,
        default_access_by_role=DEFAULT_ACCESS_LEVEL_BY_ROLE,
        today=date.today().isoformat(),
    )


@app.route("/staff/logout")
def staff_logout():
    session.pop("staff_role", None)
    session.pop("staff_id", None)
    return redirect(url_for("dashboard"))


@app.route("/theme/toggle", methods=["POST"])
def toggle_theme():
    """Any signed-in staff member can flip their own department's theme."""
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))

    new_theme = "dark" if theme_for_role(role) == "light" else "light"
    update_department_themes({role: new_theme})
    return redirect(request.referrer or url_for("workspace"))


def _build_calendar_widget(year: int, month: int):
    counts = appointment_counts_for_month(year, month)
    weeks = month_grid(year, month)
    today = date.today()
    for week in weeks:
        for cell in week:
            if cell["day"] == 0:
                cell["in_month"] = False
            else:
                cell_date = date(year, month, cell["day"])
                cell["in_month"] = True
                cell["date"] = cell_date.isoformat()
                cell["count"] = counts.get(cell["day"], 0)
                cell["is_today"] = cell_date == today

    prev_month, prev_year = (12, year - 1) if month == 1 else (month - 1, year)
    next_month, next_year = (1, year + 1) if month == 12 else (month + 1, year)

    return {
        "weeks": weeks,
        "label": date(year, month, 1).strftime("%B %Y"),
        "prev_year": prev_year,
        "prev_month": prev_month,
        "next_year": next_year,
        "next_month": next_month,
    }


@app.route("/admin/schedule/<date_str>")
def admin_schedule(date_str):
    redirect_response = require_admin("patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()

    try:
        day = datetime.strptime(date_str, "%Y-%m-%d").date()
    except ValueError:
        flash("Invalid date.", "error")
        return redirect(url_for("workspace"))

    return render_template("admin_schedule.html", staff=staff, day=day, patients=appointments_on_date(day))


@app.route("/doctor/schedule/<date_str>")
def doctor_schedule(date_str):
    redirect_response = require_doctor("patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()

    try:
        day = datetime.strptime(date_str, "%Y-%m-%d").date()
    except ValueError:
        flash("Invalid date.", "error")
        return redirect(url_for("workspace"))

    return render_template("doctor_schedule.html", staff=staff, day=day, patients=appointments_on_date(day))


@app.route("/dashboard/layout/reorder", methods=["POST"])
def dashboard_reorder():
    role, staff = current_staff()
    if not staff or role not in DEFAULT_DASHBOARD_WIDGETS:
        return {"error": "forbidden"}, 403

    data = request.get_json(silent=True) or {}
    order = data.get("order", [])
    if not isinstance(order, list):
        return {"error": "invalid order"}, 400

    save_dashboard_order(role, staff.id, order)
    return {"ok": True}


@app.route("/dashboard/widget/<widget_key>/toggle", methods=["POST"])
def toggle_dashboard_widget_route(widget_key):
    role, staff = current_staff()
    if not staff or role not in DEFAULT_DASHBOARD_WIDGETS:
        return redirect(url_for("staff_login"))

    toggle_dashboard_widget(role, staff.id, widget_key)
    return redirect(url_for("workspace"))


@app.route("/workspace")
def workspace():
    if not current_user():
        return redirect(url_for("clinic_login"))
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))

    metrics = None
    charts = None
    calendar_widget = None
    layout = None

    if role == "admin":
        metrics = {
            "total_patients": total_patient_count(),
            "new_patients_this_week": new_patients_since(datetime.now() - timedelta(days=7)),
            "total_staff": total_staff_count(),
            "staff_counts": staff_counts_by_role(),
            "open_tickets": open_ticket_count(),
            "overdue_appointments": overdue_appointment_count(),
        }

        per_day = new_patients_per_day(7)
        growth = patient_growth(30)
        staff_counts = staff_counts_by_role()

        charts = {
            "new_patients_bar": bar_chart(
                values=[count for _, count in per_day],
                labels=[day.strftime("%a") for day, _ in per_day],
            ),
            "patient_growth_line": line_chart(values=[count for _, count in growth]),
            "staff_pie": pie_chart({STAFF_LABELS[role_key]: count for role_key, count in staff_counts.items()}),
        }
    elif role == "doctor":
        metrics = {
            "total_patients": total_patient_count(),
            "my_consultations": count_consultations_by_doctor(staff.id),
            "my_prescriptions": count_prescriptions_by_doctor(staff.id),
            "unread_notes": unread_notes_count("doctor", staff.id),
        }

    if role in DEFAULT_DASHBOARD_WIDGETS:
        today = date.today()
        cal_year = request.args.get("cal_year", type=int) or today.year
        cal_month = request.args.get("cal_month", type=int) or today.month
        calendar_widget = _build_calendar_widget(cal_year, cal_month)
        layout = get_dashboard_layout(role, staff.id)

    return render_template(
        "workspace.html",
        role=role,
        role_label=STAFF_LABELS[role],
        staff=staff,
        staff_labels=STAFF_LABELS,
        metrics=metrics,
        charts=charts,
        calendar_widget=calendar_widget,
        layout=layout,
    )


@app.route("/search")
def search():
    if not current_user():
        return redirect(url_for("clinic_login"))
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))

    q = request.args.get("q", "").strip()

    patients = list_patients(q) if q else []
    staff_matches = search_staff(q) if q and role == "admin" else []
    pages = [
        (label, endpoint)
        for label, endpoint in SEARCHABLE_PAGES.get(role, [])
        if q and q.lower() in label.lower()
    ]

    return render_template(
        "search_results.html",
        q=q,
        patients=patients,
        staff_matches=staff_matches,
        staff_labels=STAFF_LABELS,
        pages=pages,
        role=role,
        role_label=STAFF_LABELS[role],
        staff=staff,
    )


# -------------------- Help / support --------------------

@app.route("/help", methods=["GET", "POST"])
def help_center():
    if not current_user():
        return redirect(url_for("clinic_login"))
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))

    if request.method == "POST":
        subject = request.form.get("subject", "").strip()
        message = request.form.get("message", "").strip()

        errors = []
        if not subject:
            errors.append("Enter a subject.")
        if not message:
            errors.append("Describe your issue or question.")

        if errors:
            for e in errors:
                flash(e, "error")
            return redirect(url_for("help_center"))

        create_ticket(role, staff.full_name, subject, message)
        flash("Support ticket submitted.", "success")
        return redirect(url_for("help_center"))

    return render_template(
        "help.html",
        role=role,
        role_label=STAFF_LABELS[role],
        staff=staff,
        faq_items=FAQ_ITEMS,
        my_tickets=list_tickets_by_submitter(role, staff.full_name),
    )


@app.route("/admin/tickets")
def admin_tickets():
    redirect_response = require_admin("admin_ops", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    return render_template("admin_tickets.html", staff=staff, tickets=list_all_tickets(), staff_labels=STAFF_LABELS)


@app.route("/admin/tickets/<int:ticket_id>/resolve", methods=["POST"])
def admin_resolve_ticket(ticket_id):
    redirect_response = require_admin("admin_ops", "write")
    if redirect_response:
        return redirect_response

    ticket = get_ticket(ticket_id)
    if ticket is None:
        flash("Ticket not found.", "error")
        return redirect(url_for("admin_tickets"))

    resolve_ticket(ticket)
    flash("Ticket marked resolved.", "success")
    return redirect(url_for("admin_tickets"))


# -------------------- Admin: patient management --------------------

def require_admin(module="admin_ops", need="view"):
    """Returns None if the caller is a signed-in admin with the required
    access, else a redirect. `module`/`need` are checked via has_access()
    against the admin's access_level."""
    if not current_user():
        return redirect(url_for("clinic_login"))
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))
    if role != "admin":
        return redirect(url_for("workspace"))
    if not has_access(staff, module, need):
        flash("Your account doesn't have access to that.", "error")
        return redirect(url_for("workspace"))
    return None


PER_PAGE_OPTIONS = (10, 25, 50, 100)


def _patients_for_list(q):
    """Reads ?sort=&dir=, the filter params, and ?page=&per_page= from the
    request. Returns (patients, last_visits, sort, direction, filters,
    pagination, page_args). 'last_visit' isn't a DB column (it's computed
    from consultations), so that sort case runs in Python after fetching,
    and pagination slices the already-sorted list rather than the query
    itself, for the same reason."""
    sort = request.args.get("sort", "created_at")
    direction = request.args.get("dir", "desc")
    if direction not in ("asc", "desc"):
        direction = "desc"

    def _parse_int(value):
        return int(value) if value.isdigit() else None

    def _parse_date(value):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return None

    filters = {
        "gender": request.args.get("gender", "").strip(),
        "status": request.args.get("status", "").strip(),
        "min_age": request.args.get("min_age", "").strip(),
        "max_age": request.args.get("max_age", "").strip(),
        "registered_from": request.args.get("registered_from", "").strip(),
        "registered_to": request.args.get("registered_to", "").strip(),
        "location": request.args.get("location", "").strip(),
    }
    filter_kwargs = dict(
        gender=filters["gender"],
        status=filters["status"],
        min_age=_parse_int(filters["min_age"]),
        max_age=_parse_int(filters["max_age"]),
        registered_from=_parse_date(filters["registered_from"]),
        registered_to=_parse_date(filters["registered_to"]),
        location=filters["location"],
    )
    filters["active_count"] = sum(1 for v in filters.values() if v)

    if sort == "last_visit":
        patients = list_patients(q, **filter_kwargs)
        last_visits = {patient.id: last_visit(patient.id) for patient in patients}
        patients.sort(key=lambda p: last_visits[p.id] or datetime.min, reverse=(direction == "desc"))
    else:
        patients = list_patients(q, sort=sort, direction=direction, **filter_kwargs)
        last_visits = {patient.id: last_visit(patient.id) for patient in patients}

    per_page_raw = request.args.get("per_page", "25")
    per_page = int(per_page_raw) if per_page_raw.isdigit() and int(per_page_raw) in PER_PAGE_OPTIONS else 25

    page_raw = request.args.get("page", "1")
    page = int(page_raw) if page_raw.isdigit() and int(page_raw) >= 1 else 1

    total = len(patients)
    total_pages = max(1, -(-total // per_page))
    page = min(page, total_pages)

    start_idx = (page - 1) * per_page
    end_idx = start_idx + per_page
    page_patients = patients[start_idx:end_idx]
    page_last_visits = {p.id: last_visits[p.id] for p in page_patients}

    pagination = {
        "page": page,
        "per_page": per_page,
        "per_page_options": PER_PAGE_OPTIONS,
        "total": total,
        "total_pages": total_pages,
        "has_prev": page > 1,
        "has_next": page < total_pages,
        "start": start_idx + 1 if total else 0,
        "end": min(end_idx, total),
    }

    page_args = {
        "q": q,
        "sort": sort,
        "dir": direction,
        "per_page": per_page,
        "gender": filters["gender"],
        "status": filters["status"],
        "min_age": filters["min_age"],
        "max_age": filters["max_age"],
        "registered_from": filters["registered_from"],
        "registered_to": filters["registered_to"],
        "location": filters["location"],
    }

    return page_patients, page_last_visits, sort, direction, filters, pagination, page_args


def _validate_patient_form(form):
    fields = {
        "full_name": form.get("full_name", "").strip(),
        "gender": form.get("gender", "").strip(),
        "phone_number": form.get("phone_number", "").strip(),
        "address": form.get("address", "").strip(),
        "email": form.get("email", "").strip(),
        "status": form.get("status", "Active").strip(),
        "insurance_provider": form.get("insurance_provider", "").strip(),
        "insurance_policy_number": form.get("insurance_policy_number", "").strip(),
        "emergency_contact_name": form.get("emergency_contact_name", "").strip(),
        "emergency_contact_phone": form.get("emergency_contact_phone", "").strip(),
    }

    errors = []
    if not fields["full_name"]:
        errors.append("Enter the patient's full name.")

    date_of_birth_raw = form.get("date_of_birth", "").strip()
    fields["date_of_birth"] = None
    if not date_of_birth_raw:
        errors.append("Enter the patient's date of birth.")
    else:
        try:
            fields["date_of_birth"] = datetime.strptime(date_of_birth_raw, "%Y-%m-%d").date()
        except ValueError:
            errors.append("Date of birth is not a valid date.")

    next_appointment_raw = form.get("next_appointment", "").strip()
    fields["next_appointment"] = None
    if next_appointment_raw:
        try:
            fields["next_appointment"] = datetime.strptime(next_appointment_raw, "%Y-%m-%d").date()
        except ValueError:
            errors.append("Next appointment is not a valid date.")

    if fields["gender"] not in ("Female", "Male", "Other"):
        errors.append("Choose a gender.")
    if not fields["phone_number"]:
        errors.append("Enter a phone number.")
    if fields["status"] not in PATIENT_STATUSES:
        errors.append("Choose a valid status.")

    return errors, fields


@app.route("/admin/patients")
def admin_patients():
    redirect_response = require_admin("patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    q = request.args.get("q", "").strip()
    patients, last_visits, sort, direction, filters, pagination, page_args = _patients_for_list(q)
    missing_fields = {patient.id: patient_missing_fields(patient) for patient in patients}
    overdue = {patient.id: is_overdue(patient) for patient in patients}
    return render_template(
        "admin_patients.html",
        patients=patients,
        last_visits=last_visits,
        q=q,
        staff=staff,
        sort=sort,
        dir=direction,
        filters=filters,
        statuses=PATIENT_STATUSES,
        pagination=pagination,
        page_args=page_args,
        missing_fields=missing_fields,
        overdue=overdue,
    )


@app.route("/admin/patients/<int:patient_id>")
def admin_patient_profile(patient_id):
    redirect_response = require_admin("patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()

    patient = get_patient(patient_id)
    if patient is None:
        flash("Patient not found.", "error")
        return redirect(url_for("admin_patients"))

    return render_template(
        "admin_patient_profile.html",
        patient=patient,
        staff=staff,
        last_visit=last_visit(patient.id),
        consultations=patient_consultations(patient.id),
        prescriptions=patient_prescriptions(patient.id),
        lab_results=patient_lab_results(patient.id),
        missing_fields=patient_missing_fields(patient),
        overdue=is_overdue(patient),
    )


@app.route("/admin/patients/register", methods=["GET", "POST"])
def admin_register_patient():
    redirect_response = require_admin("patients", "write" if request.method == "POST" else "view")
    if redirect_response:
        return redirect_response

    if request.method == "POST":
        errors, fields = _validate_patient_form(request.form)

        patient_number_mode = request.form.get("patient_number_mode", "auto")
        manual_number = request.form.get("patient_number", "").strip()
        if patient_number_mode == "manual":
            if not manual_number.isdigit() or len(manual_number) != 6:
                errors.append("Patient number must be exactly 6 digits.")
            elif patient_number_taken(manual_number):
                errors.append("That patient number is already registered.")
        else:
            manual_number = None

        if errors:
            for e in errors:
                flash(e, "error")
            return redirect(url_for("admin_patients", open_register=1))

        duplicates = find_potential_duplicate_patients(fields["full_name"], fields["date_of_birth"], fields["phone_number"])
        patient = create_patient(**fields, patient_number=manual_number)

        if duplicates:
            existing = ", ".join(f"{d.full_name} (#{d.patient_number})" for d in duplicates)
            flash(f"Possible duplicate: {existing} has a matching name/DOB or phone number.", "warning")
        flash(f"Patient registered. Patient number: {patient.patient_number}.", "success")
        return redirect(url_for("admin_patients"))

    # The registration form is now a modal on the patients list, not its own page.
    return redirect(url_for("admin_patients", open_register=1))


@app.route("/admin/patients/<int:patient_id>/edit", methods=["GET", "POST"])
def admin_edit_patient(patient_id):
    redirect_response = require_admin("patients", "write" if request.method == "POST" else "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()

    patient = get_patient(patient_id)
    if patient is None:
        flash("Patient not found.", "error")
        return redirect(url_for("admin_patients"))

    if request.method == "POST":
        errors, fields = _validate_patient_form(request.form)

        if errors:
            for e in errors:
                flash(e, "error")
            return redirect(url_for("admin_edit_patient", patient_id=patient_id))

        update_patient(patient, **fields)
        flash("Patient updated.", "success")
        return redirect(url_for("admin_patients"))

    return render_template("admin_edit_patient.html", patient=patient, staff=staff, statuses=PATIENT_STATUSES)


@app.route("/admin/patients/<int:patient_id>/delete", methods=["POST"])
def admin_delete_patient(patient_id):
    redirect_response = require_admin("patients", "write")
    if redirect_response:
        return redirect_response

    patient = get_patient(patient_id)
    if patient is None:
        flash("Patient not found.", "error")
        return redirect(url_for("admin_patients"))

    delete_patient(patient)
    flash(f"{patient.full_name} was removed from patient records.", "success")
    return redirect(url_for("admin_patients"))


def _report_data():
    return {
        "total_patients": total_patient_count(),
        "staff_counts": staff_counts_by_role(),
        "total_staff": total_staff_count(),
        "recent_patients": recent_patients(),
        "recent_staff": recent_staff(),
        "demographics": demographics_summary(),
        "age_gender_rows": age_gender_breakdown(),
        "registrations_by_month": new_registrations_by_month(12),
        "generated_at": datetime.now(),
    }


@app.route("/admin/report")
def admin_report():
    redirect_response = require_admin("admin_ops", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    return render_template("admin_report.html", staff_labels=STAFF_LABELS, staff=staff, **_report_data())


@app.route("/admin/report/print")
def admin_report_print():
    redirect_response = require_admin("admin_ops", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    return render_template("admin_report_print.html", staff_labels=STAFF_LABELS, staff=staff, **_report_data())


@app.route("/admin/report/export.csv")
def admin_report_export_csv():
    redirect_response = require_admin("admin_ops", "view")
    if redirect_response:
        return redirect_response
    data = _report_data()

    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow(["Demographics Summary"])
    writer.writerow(["Metric", "Value"])
    writer.writerow(["Total Patients", data["demographics"]["total"]])
    for gender, count in data["demographics"]["gender_counts"].items():
        writer.writerow([gender, count])
    for status, count in data["demographics"]["status_counts"].items():
        writer.writerow([status, count])
    writer.writerow([])

    writer.writerow(["Age / Gender Breakdown"])
    writer.writerow(["Age Group", "Female", "Male", "Other", "Total"])
    for row in data["age_gender_rows"]:
        writer.writerow([row["age_group"], row["Female"], row["Male"], row["Other"], row["total"]])
    writer.writerow([])

    writer.writerow(["New Registrations Over Time (last 12 months)"])
    writer.writerow(["Month", "New Patients"])
    for month_label, count in data["registrations_by_month"]:
        writer.writerow([month_label, count])

    response = Response(buf.getvalue(), mimetype="text/csv")
    response.headers["Content-Disposition"] = "attachment; filename=clinic_report.csv"
    return response


@app.route("/admin/report/export.xlsx")
def admin_report_export_xlsx():
    redirect_response = require_admin("admin_ops", "view")
    if redirect_response:
        return redirect_response
    data = _report_data()

    workbook = Workbook()

    demographics_sheet = workbook.active
    demographics_sheet.title = "Demographics"
    demographics_sheet.append(["Metric", "Value"])
    demographics_sheet.append(["Total Patients", data["demographics"]["total"]])
    for gender, count in data["demographics"]["gender_counts"].items():
        demographics_sheet.append([gender, count])
    for status, count in data["demographics"]["status_counts"].items():
        demographics_sheet.append([status, count])

    breakdown_sheet = workbook.create_sheet("Age-Gender Breakdown")
    breakdown_sheet.append(["Age Group", "Female", "Male", "Other", "Total"])
    for row in data["age_gender_rows"]:
        breakdown_sheet.append([row["age_group"], row["Female"], row["Male"], row["Other"], row["total"]])

    registrations_sheet = workbook.create_sheet("Registrations Over Time")
    registrations_sheet.append(["Month", "New Patients"])
    for month_label, count in data["registrations_by_month"]:
        registrations_sheet.append([month_label, count])

    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)

    response = Response(
        buf.read(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response.headers["Content-Disposition"] = "attachment; filename=clinic_report.xlsx"
    return response


@app.route("/admin/settings")
def admin_settings():
    redirect_response = require_admin("admin_ops", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    return render_template(
        "admin_settings.html",
        user=current_user(),
        settings=get_settings(),
        staff=staff,
        staff_labels=STAFF_LABELS,
        all_staff=Staff.query.order_by(Staff.role, Staff.full_name).all(),
        access_level_labels=ACCESS_LEVEL_LABELS,
    )


@app.route("/admin/settings/profile", methods=["POST"])
def admin_settings_profile():
    redirect_response = require_admin("admin_ops", "write")
    if redirect_response:
        return redirect_response

    hospital_name = request.form.get("hospital_name", "").strip()
    address = request.form.get("address", "").strip()
    phone_number = request.form.get("phone_number", "").strip()
    operating_hours = request.form.get("operating_hours", "").strip()

    if not hospital_name:
        flash("Enter the hospital name.", "error")
        return redirect(url_for("admin_settings"))

    user = current_user()
    user.hospital_name = hospital_name
    db.session.commit()
    update_clinic_profile(address, phone_number, operating_hours)
    flash("Clinic profile updated.", "success")
    return redirect(url_for("admin_settings"))


@app.route("/admin/settings/pin-policy", methods=["POST"])
def admin_settings_pin_policy():
    redirect_response = require_admin("admin_ops", "write")
    if redirect_response:
        return redirect_response

    pin_length_raw = request.form.get("pin_length", "").strip()
    if not pin_length_raw.isdigit() or not (4 <= int(pin_length_raw) <= 8):
        flash("PIN length must be a number between 4 and 8.", "error")
        return redirect(url_for("admin_settings"))

    update_pin_policy(int(pin_length_raw))
    flash("PIN policy updated. This applies to staff added from now on.", "success")
    return redirect(url_for("admin_settings"))


@app.route("/admin/settings/announcement", methods=["POST"])
def admin_settings_announcement():
    redirect_response = require_admin("admin_ops", "write")
    if redirect_response:
        return redirect_response

    text = request.form.get("announcement", "").strip()
    update_announcement(text)
    flash("Announcement updated." if text else "Announcement cleared.", "success")
    return redirect(url_for("admin_settings"))


@app.route("/admin/settings/themes", methods=["POST"])
def admin_settings_themes():
    redirect_response = require_admin("admin_ops", "write")
    if redirect_response:
        return redirect_response

    themes = {}
    for role in STAFF_ROLES:
        theme = request.form.get(f"{role}_theme", "")
        if theme not in ("light", "dark"):
            flash("Choose light or dark for every department.", "error")
            return redirect(url_for("admin_settings"))
        themes[role] = theme

    update_department_themes(themes)
    flash("Department themes updated.", "success")
    return redirect(url_for("admin_settings"))


@app.route("/admin/settings/password", methods=["POST"])
def admin_settings_password():
    redirect_response = require_admin("admin_ops", "write")
    if redirect_response:
        return redirect_response

    current_password = request.form.get("current_password", "")
    new_password = request.form.get("new_password", "")
    confirm = request.form.get("confirm", "")

    user = current_user()
    errors = []
    if not user.check_password(current_password):
        errors.append("Current password is incorrect.")
    if len(new_password) < 6:
        errors.append("New password must be at least 6 characters.")
    if new_password != confirm:
        errors.append("New passwords don't match.")

    if errors:
        for e in errors:
            flash(e, "error")
        return redirect(url_for("admin_settings"))

    user.set_password(new_password)
    db.session.commit()
    flash("Password changed.", "success")
    return redirect(url_for("admin_settings"))


# -------------------- Doctor: patients, consultations, prescriptions --------------------

def require_doctor(module="patients", need="view"):
    """Returns None if the caller is a signed-in doctor with the required
    access, else a redirect."""
    if not current_user():
        return redirect(url_for("clinic_login"))
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))
    if role != "doctor":
        return redirect(url_for("workspace"))
    if not has_access(staff, module, need):
        flash("Your account doesn't have access to that.", "error")
        return redirect(url_for("workspace"))
    return None


def require_role(role_name, module="patients", need="view"):
    """Generic guard for a single-role, single-page area (Receptionist, Lab
    Technician). Returns None if allowed, else a redirect."""
    if not current_user():
        return redirect(url_for("clinic_login"))
    role, staff = current_staff()
    if not staff:
        return redirect(url_for("staff_login"))
    if role != role_name:
        return redirect(url_for("workspace"))
    if not has_access(staff, module, need):
        flash("Your account doesn't have access to that.", "error")
        return redirect(url_for("workspace"))
    return None


@app.route("/doctor/patients")
def doctor_patients():
    redirect_response = require_doctor("patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    q = request.args.get("q", "").strip()
    patients, last_visits, sort, direction, filters, pagination, page_args = _patients_for_list(q)
    missing_fields = {patient.id: patient_missing_fields(patient) for patient in patients}
    overdue = {patient.id: is_overdue(patient) for patient in patients}
    unread_by_patient = unread_notes_by_patient("doctor", staff.id)
    return render_template(
        "doctor_patients.html",
        patients=patients,
        last_visits=last_visits,
        q=q,
        staff=staff,
        sort=sort,
        dir=direction,
        filters=filters,
        statuses=PATIENT_STATUSES,
        pagination=pagination,
        page_args=page_args,
        missing_fields=missing_fields,
        overdue=overdue,
        unread_by_patient=unread_by_patient,
    )


@app.route("/doctor/patients/<int:patient_id>", methods=["GET", "POST"])
def doctor_patient_detail(patient_id):
    redirect_response = require_doctor("patients", "write" if request.method == "POST" else "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()

    patient = get_patient(patient_id)
    if patient is None:
        flash("Patient not found.", "error")
        return redirect(url_for("doctor_patients"))

    if request.method == "POST":
        form_name = request.form.get("form_name", "")

        if form_name == "consultation":
            notes = request.form.get("notes", "").strip()
            if not notes:
                flash("Enter consultation notes.", "error")
            else:
                add_consultation(patient.id, staff.id, notes)
                flash("Consultation note added.", "success")

        elif form_name == "prescription":
            medication_name = request.form.get("medication_name", "").strip()
            dosage = request.form.get("dosage", "").strip()
            instructions = request.form.get("instructions", "").strip()
            errors = []
            if not medication_name:
                errors.append("Enter the medication name.")
            if not dosage:
                errors.append("Enter the dosage.")
            if not instructions:
                errors.append("Enter instructions.")
            if errors:
                for e in errors:
                    flash(e, "error")
            else:
                add_prescription(patient.id, staff.id, medication_name, dosage, instructions)
                flash("Prescription added.", "success")

        return redirect(url_for("doctor_patient_detail", patient_id=patient.id))

    mark_patient_notes_read("doctor", staff.id, patient.id)

    return render_template(
        "doctor_patient_detail.html",
        patient=patient,
        staff=staff,
        consultations=patient_consultations(patient.id),
        prescriptions=patient_prescriptions(patient.id),
        lab_results=patient_lab_results(patient.id),
        missing_fields=patient_missing_fields(patient),
        overdue=is_overdue(patient),
    )


# -------------------- Receptionist: view/search patients only --------------------

@app.route("/receptionist/patients")
def receptionist_patients():
    redirect_response = require_role("receptionist", "patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    q = request.args.get("q", "").strip()
    patients, last_visits, sort, direction, filters, pagination, page_args = _patients_for_list(q)
    missing_fields = {patient.id: patient_missing_fields(patient) for patient in patients}
    overdue = {patient.id: is_overdue(patient) for patient in patients}
    return render_template(
        "receptionist_patients.html",
        patients=patients,
        last_visits=last_visits,
        q=q,
        staff=staff,
        sort=sort,
        dir=direction,
        filters=filters,
        statuses=PATIENT_STATUSES,
        pagination=pagination,
        page_args=page_args,
        missing_fields=missing_fields,
        overdue=overdue,
    )


# -------------------- Lab Technician: record & view lab results --------------------

@app.route("/lab/patients")
def lab_patients():
    redirect_response = require_role("lab_technician", "patients", "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()
    q = request.args.get("q", "").strip()
    patients, last_visits, sort, direction, filters, pagination, page_args = _patients_for_list(q)
    return render_template(
        "lab_patients.html",
        patients=patients,
        last_visits=last_visits,
        q=q,
        staff=staff,
        sort=sort,
        dir=direction,
        filters=filters,
        statuses=PATIENT_STATUSES,
        pagination=pagination,
        page_args=page_args,
    )


@app.route("/lab/patients/<int:patient_id>", methods=["GET", "POST"])
def lab_patient_detail(patient_id):
    redirect_response = require_role("lab_technician", "patients", "write" if request.method == "POST" else "view")
    if redirect_response:
        return redirect_response
    _, staff = current_staff()

    patient = get_patient(patient_id)
    if patient is None:
        flash("Patient not found.", "error")
        return redirect(url_for("lab_patients"))

    if request.method == "POST":
        test_name = request.form.get("test_name", "").strip()
        result = request.form.get("result", "").strip()
        notes = request.form.get("notes", "").strip()
        errors = []
        if not test_name:
            errors.append("Enter the test name.")
        if not result:
            errors.append("Enter the result.")
        if errors:
            for e in errors:
                flash(e, "error")
        else:
            add_lab_result(patient.id, staff.id, test_name, result, notes)
            flash("Lab result recorded.", "success")
        return redirect(url_for("lab_patient_detail", patient_id=patient.id))

    return render_template(
        "lab_patient_detail.html",
        patient=patient,
        staff=staff,
        lab_results=patient_lab_results(patient.id),
    )


# -------------------- Dashboard --------------------

@app.route("/dashboard")
def dashboard():
    user = current_user()
    if not user:
        return redirect(url_for("clinic_login"))
    return render_template("dashboard.html", user=user)


if __name__ == "__main__":
    app.run(debug=True)
